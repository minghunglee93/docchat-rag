"""
Document Processing Module
Handles loading, parsing, and chunking documents
"""

from typing import List, Dict, Optional
from pathlib import Path
import re
from langchain_text_splitters import RecursiveCharacterTextSplitter
from langchain_core.documents import Document
import pypdf
import docx
import requests
from bs4 import BeautifulSoup


class DocumentProcessor:
    """Process and chunk documents for RAG"""

    def __init__(
            self,
            chunk_size: int = 1000,
            chunk_overlap: int = 200,
            separators: Optional[List[str]] = None
    ):
        """
        Initialize document processor

        Args:
            chunk_size: Size of each text chunk
            chunk_overlap: Overlap between chunks
            separators: Custom text separators
        """
        self.chunk_size = chunk_size
        self.chunk_overlap = chunk_overlap

        # Default separators for better chunking
        if separators is None:
            separators = ["\n\n", "\n", ". ", " ", ""]

        self.text_splitter = RecursiveCharacterTextSplitter(
            chunk_size=chunk_size,
            chunk_overlap=chunk_overlap,
            separators=separators,
            length_function=len
        )

    def load_pdf(self, file_path: str) -> List[Document]:
        """
        Load and parse PDF file

        Args:
            file_path: Path to PDF file

        Returns:
            List of Document objects
        """
        documents = []

        try:
            pdf_reader = pypdf.PdfReader(file_path)

            for page_num, page in enumerate(pdf_reader.pages):
                text = page.extract_text()

                if text.strip():
                    doc = Document(
                        page_content=text,
                        metadata={
                            "source": file_path,
                            "page": page_num + 1,
                            "total_pages": len(pdf_reader.pages)
                        }
                    )
                    documents.append(doc)

            print(f"✓ Loaded PDF: {len(documents)} pages from {Path(file_path).name}")

        except Exception as e:
            print(f"Error loading PDF {file_path}: {e}")
            raise

        return documents

    def load_docx(self, file_path: str) -> List[Document]:
        """
        Load and parse DOCX file

        Args:
            file_path: Path to DOCX file

        Returns:
            List of Document objects
        """
        documents = []

        try:
            doc = docx.Document(file_path)

            # Combine all paragraphs
            text = "\n\n".join([para.text for para in doc.paragraphs if para.text.strip()])

            if text.strip():
                document = Document(
                    page_content=text,
                    metadata={
                        "source": file_path,
                        "paragraphs": len(doc.paragraphs)
                    }
                )
                documents.append(document)

            print(f"✓ Loaded DOCX: {Path(file_path).name}")

        except Exception as e:
            print(f"Error loading DOCX {file_path}: {e}")
            raise

        return documents

    def load_txt(self, file_path: str) -> List[Document]:
        """
        Load and parse TXT file

        Args:
            file_path: Path to TXT file

        Returns:
            List of Document objects
        """
        documents = []

        try:
            with open(file_path, 'r', encoding='utf-8') as f:
                text = f.read()

            if text.strip():
                document = Document(
                    page_content=text,
                    metadata={
                        "source": file_path,
                        "size": len(text)
                    }
                )
                documents.append(document)

            print(f"✓ Loaded TXT: {Path(file_path).name}")

        except Exception as e:
            print(f"Error loading TXT {file_path}: {e}")
            raise

        return documents

    def load_url(self, url: str) -> List[Document]:
        """
        Load and parse web page

        Args:
            url: URL to fetch

        Returns:
            List of Document objects
        """
        documents = []

        try:
            response = requests.get(url, timeout=10)
            response.raise_for_status()

            soup = BeautifulSoup(response.content, 'html.parser')

            # Remove script and style elements
            for script in soup(["script", "style"]):
                script.decompose()

            # Get text
            text = soup.get_text()

            # Clean up text
            lines = (line.strip() for line in text.splitlines())
            chunks = (phrase.strip() for line in lines for phrase in line.split("  "))
            text = '\n'.join(chunk for chunk in chunks if chunk)

            if text.strip():
                document = Document(
                    page_content=text,
                    metadata={
                        "source": url,
                        "type": "webpage"
                    }
                )
                documents.append(document)

            print(f"✓ Loaded URL: {url}")

        except Exception as e:
            print(f"Error loading URL {url}: {e}")
            raise

        return documents

    def load_document(self, file_path: str) -> List[Document]:
        """
        Load document based on file extension

        Args:
            file_path: Path to file or URL

        Returns:
            List of Document objects
        """
        # Check if URL
        if file_path.startswith('http://') or file_path.startswith('https://'):
            return self.load_url(file_path)

        # Check file extension
        path = Path(file_path)
        extension = path.suffix.lower()

        if extension == '.pdf':
            return self.load_pdf(file_path)
        elif extension == '.docx':
            return self.load_docx(file_path)
        elif extension == '.txt':
            return self.load_txt(file_path)
        else:
            raise ValueError(f"Unsupported file type: {extension}")

    def chunk_documents(self, documents: List[Document]) -> List[Document]:
        """
        Split documents into chunks

        Args:
            documents: List of Document objects

        Returns:
            List of chunked Document objects
        """
        chunks = []

        for doc in documents:
            # Split document into chunks
            splits = self.text_splitter.split_text(doc.page_content)

            # Create new Document for each chunk
            for i, split in enumerate(splits):
                chunk = Document(
                    page_content=split,
                    metadata={
                        **doc.metadata,
                        "chunk": i,
                        "total_chunks": len(splits)
                    }
                )
                chunks.append(chunk)

        print(f"✓ Created {len(chunks)} chunks from {len(documents)} documents")

        return chunks

    def process_file(self, file_path: str) -> List[Document]:
        """
        Complete pipeline: load and chunk document

        Args:
            file_path: Path to file or URL

        Returns:
            List of chunked Document objects
        """
        # Load document
        documents = self.load_document(file_path)

        # Chunk documents
        chunks = self.chunk_documents(documents)

        return chunks

    def process_multiple_files(self, file_paths: List[str]) -> List[Document]:
        """
        Process multiple files

        Args:
            file_paths: List of file paths or URLs

        Returns:
            List of chunked Document objects
        """
        all_chunks = []

        print(f"\nProcessing {len(file_paths)} files...")
        print("=" * 60)

        for file_path in file_paths:
            try:
                chunks = self.process_file(file_path)
                all_chunks.extend(chunks)
            except Exception as e:
                print(f"✗ Failed to process {file_path}: {e}")

        print("=" * 60)
        print(f"✓ Total chunks: {len(all_chunks)}")

        return all_chunks

    def get_document_stats(self, documents: List[Document]) -> Dict:
        """
        Get statistics about documents

        Args:
            documents: List of Document objects

        Returns:
            Dictionary with statistics
        """
        total_chars = sum(len(doc.page_content) for doc in documents)
        total_words = sum(len(doc.page_content.split()) for doc in documents)

        sources = set()
        for doc in documents:
            if 'source' in doc.metadata:
                sources.add(doc.metadata['source'])

        stats = {
            "total_documents": len(documents),
            "total_characters": total_chars,
            "total_words": total_words,
            "avg_chars_per_doc": total_chars / len(documents) if documents else 0,
            "avg_words_per_doc": total_words / len(documents) if documents else 0,
            "unique_sources": len(sources)
        }

        return stats


# Example usage and testing
if __name__ == "__main__":
    import sys

    # Initialize processor
    processor = DocumentProcessor(chunk_size=1000, chunk_overlap=200)

    if len(sys.argv) > 1:
        # Process files from command line
        file_paths = sys.argv[1:]
        chunks = processor.process_multiple_files(file_paths)

        # Show statistics
        stats = processor.get_document_stats(chunks)
        print("\nDocument Statistics:")
        print("=" * 60)
        for key, value in stats.items():
            print(f"{key}: {value}")

        # Show sample chunks
        print("\nSample Chunks:")
        print("=" * 60)
        for i, chunk in enumerate(chunks[:3], 1):
            print(f"\nChunk {i}:")
            print(f"Source: {chunk.metadata.get('source', 'Unknown')}")
            print(f"Content: {chunk.page_content[:200]}...")
            print("-" * 60)
    else:
        print("Usage: python document_processor.py <file1> <file2> ...")
        print("\nSupported formats:")
        print("  - PDF files (.pdf)")
        print("  - Word documents (.docx)")
        print("  - Text files (.txt)")
        print("  - URLs (http:// or https://)")